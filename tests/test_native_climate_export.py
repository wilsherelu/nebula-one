"""Real HTTP acceptance of the native climate artifacts, on synthetic evidence.

The climate answer these tests export is the one the pinned contract fixture returned, so the
artifact route is exercised through the persisted column rather than a re-solve.
"""

from __future__ import annotations

import copy
import io
import json

from openpyxl import load_workbook
import pytest

from nebula_one.runtime_models import SolvedSystemSnapshot

from nebula_one.exports.native_climate import build_native_climate_html, build_native_climate_xlsx
from nebula_one.mapped_background_schemas import (
    ProviderNativeBackgroundClimateRequest,
    ProviderNativeBackgroundClimateResponse,
)
from nebula_one.native_climate_contract import NativeClimateContract

import test_native_climate_solve as native_solve

# The solve fixture is owned by the native solve acceptance module.
native_setup = native_solve.native_setup

FACTOR_SOURCE = "ecoinvent_3.11_native_EF_v3.1_climate_total"
PER_EXCHANGE_MARKERS = (
    "synthetic::carbon-storage-in",
    "synthetic::uncharacterised-out",
    "08a91e70-3ddc-11dd-923d-0050c2490048",
    "99a91e70-3ddc-11dd-923d-0050c2490048",
    "aaa91e70-3ddc-11dd-923d-0050c2490048",
    "air, unspecified",
    "air, unfinalised",
    FACTOR_SOURCE,
)


def _cells(content: bytes) -> str:
    workbook = load_workbook(io.BytesIO(content))
    return "\n".join(
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )


def _artifacts(client, snapshot_id: str) -> tuple[str, str]:
    report = client.get(
        f"/api/v1/solved-system-snapshots/{snapshot_id}/native-background-climate/report.html"
    )
    spreadsheet = client.get(
        f"/api/v1/solved-system-snapshots/{snapshot_id}/native-background-climate/export.xlsx"
    )
    assert report.status_code == 200, report.text
    assert spreadsheet.status_code == 200, spreadsheet.text
    return report.text, _cells(spreadsheet.content)


def test_native_climate_artifacts_report_values_units_versions_and_coverage(
    client, native_setup, upstream_requests
):
    solved = native_solve._solve(native_setup)
    snapshot_id = solved["id"]

    for _ in range(2):
        report, spreadsheet = _artifacts(client, snapshot_id)
        for artifact in (report, spreadsheet):
            # Values and units the source itself published.
            assert "5.5 kg CO2-Eq" in artifact
            assert "4.0 kg CO2-Eq" in artifact
            assert "1.5 kg CO2-Eq" in artifact
            # The single indicator and its method, in the source's own words.
            assert "GWP100" in artifact
            assert "climate change" in artifact
            assert "EF v3.1" in artifact
            # Coverage statements, kept as counts rather than per-exchange rows.
            assert "有非零排放因子 83 项" in artifact
            assert "未纳入特征化 2654 项" in artifact
            # Exact version references and the contract they came from.
            assert "ecoinvent-ef31-climate-total-3.11" in artifact
            assert NativeClimateContract().request_payload()["method_package_sha256"] in artifact
            assert NativeClimateContract().request_payload()["runtime_manifest_sha256"] in artifact
            assert "70dda0cf2d79108ca4ff96a203cb61eb76a9e550288e05be133e684b4747b1dd" in artifact
            assert "92016cdbbd2abd4b1a619d0048432313f9e0cf5327cf1c98a049b5a5c1155f29" in artifact
            assert NativeClimateContract().source_contract["commit"] in artifact
            assert "docs/PROVIDER_NATIVE_CLIMATE_TOTAL_V1.md" in artifact
            # The claim limit and the basis the total is per.
            assert "不构成完整产品碳足迹" in artifact
            assert "1.0 kg" in artifact
            # Negative contributions are explained, never presented as permanent removal.
            assert "不等同于永久碳移除" in artifact

        # An export must never look like a complete product carbon footprint.
        assert "产品碳足迹报告" not in report
        assert "nebula-one-pcf" not in report

    # Neither artifact is a re-solve: the frozen answer is read back as it was stored.
    assert len(native_solve._solve_requests(upstream_requests)) == 1

    identity = solved["content_sha256"]
    report = client.get(
        f"/api/v1/solved-system-snapshots/{snapshot_id}/native-background-climate/report.html"
    )
    spreadsheet = client.get(
        f"/api/v1/solved-system-snapshots/{snapshot_id}/native-background-climate/export.xlsx"
    )
    assert report.headers["ETag"] == f'"{identity}"'
    assert report.headers["X-Nebula-Result-SHA256"] == identity
    assert spreadsheet.headers["ETag"] == f'"{identity}"'
    assert spreadsheet.headers["X-Nebula-Result-SHA256"] == identity
    assert spreadsheet.headers["content-disposition"].startswith(
        f'attachment; filename="nebula-one-native-background-climate-{identity[:12]}'
    )
    assert spreadsheet.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert report.headers["content-type"].startswith("text/html")
    assert "inline" in report.headers["content-disposition"]
    workbook = load_workbook(io.BytesIO(spreadsheet.content))
    assert workbook.sheetnames == [
        "气候核算摘要",
        "来源分域与覆盖",
        "口径与限制",
        "版本与来源证据",
    ]


def test_native_climate_artifacts_carry_no_per_exchange_factor_rows(
    client, native_setup
):
    solved = native_solve._solve(native_setup)
    frozen_answer = json.dumps(solved["native_background_climate"]["response"])
    # The refusal below is only meaningful because the licensed rows really are in the frozen answer.
    for marker in PER_EXCHANGE_MARKERS:
        assert marker in frozen_answer

    report, spreadsheet = _artifacts(client, solved["id"])
    for artifact in (report, spreadsheet):
        for marker in PER_EXCHANGE_MARKERS:
            assert marker not in artifact
        for field in ("cf_value", "cf_source", "exchange_id", "flow_uuid", "subcompartment"):
            assert field not in artifact


def test_projection_snapshot_refuses_a_native_climate_artifact(client, native_setup):
    solved = native_solve._solve(
        native_setup,
        native_background_climate=False,
        idempotency_key="synthetic-projection-solve",
    )
    assert solved["native_background_climate"] is None

    for path in ("report.html", "export.xlsx"):
        response = client.get(
            f"/api/v1/solved-system-snapshots/{solved['id']}/native-background-climate/{path}"
        )
        assert response.status_code == 422, response.text
        error = response.json()["error"]
        assert error["code"] == "NATIVE_CLIMATE_ANSWER_REQUIRED"
        assert "run the climate calculation" in error["message"]


def test_missing_snapshot_is_not_invented_as_an_empty_artifact(client):
    response = client.get(
        "/api/v1/solved-system-snapshots/does-not-exist/native-background-climate/report.html"
    )
    assert response.status_code == 404, response.text
    assert response.json()["error"]["code"] == "SOLVED_SYSTEM_SNAPSHOT_NOT_FOUND"


@pytest.mark.parametrize("damage", ["shape", "identity", "demand"])
def test_native_climate_export_rejects_damaged_frozen_contract(client, native_setup, damage):
    solved = native_solve._solve(native_setup)
    # Only the disposable synthetic test database is modified.
    with client.app.state.container.database.session() as db:
        snapshot = db.get(SolvedSystemSnapshot, solved["id"])
        if damage == "demand":
            receipt = copy.deepcopy(snapshot.provider_receipt)
            receipt["calculation_basis"].pop("demand")
            snapshot.provider_receipt = receipt
        else:
            envelope = copy.deepcopy(snapshot.native_background_climate)
            if damage == "shape":
                envelope["response"]["domains"] = {}
            else:
                envelope["response"]["method_package_version"] = "other-version"
            snapshot.native_background_climate = envelope
        db.commit()
    for name in ("report.html", "export.xlsx"):
        response = client.get(
            f"/api/v1/solved-system-snapshots/{solved['id']}/native-background-climate/{name}"
        )
        assert response.status_code == 422, response.text
        assert response.json()["error"]["code"] == "NATIVE_CLIMATE_ANSWER_INVALID"


def test_split_diagnostic_is_reported_as_domains_without_a_summed_total(client, native_setup):
    """The stub engine always merges, so the split branch is fed a real frozen answer."""
    solved = native_solve._solve(native_setup)
    envelope = solved["native_background_climate"]
    answer = ProviderNativeBackgroundClimateResponse.model_validate(
        {
            **envelope["response"],
            "merge_status": "split_diagnostic",
            "unified_climate_total": None,
            "merge_blockers": [
                {
                    "domain": "native_background",
                    "reason": "unit_group_mismatch",
                    "detail": {"expected": "mass", "actual": "energy"},
                }
            ],
        }
    )
    request = ProviderNativeBackgroundClimateRequest.model_validate(envelope["request"])
    copy_of_envelope = copy.deepcopy(envelope)

    artifacts = (
        build_native_climate_html(
            copy_of_envelope,
            answer,
            request,
            snapshot_id=solved["id"],
            snapshot_sha256=solved["content_sha256"],
            basis="1.0 kg",
        ),
        _cells(
            build_native_climate_xlsx(
                copy_of_envelope,
                answer,
                request,
                snapshot_id=solved["id"],
                snapshot_sha256=solved["content_sha256"],
                basis="1.0 kg",
            )
        ),
    )
    for artifact in artifacts:
        assert "5.5 kg CO2-Eq" not in artifact
        assert "来源未合并" in artifact
        assert "domain: native_background · reason: unit_group_mismatch · detail: " in artifact
        assert "把两域数值相加会得到一个来源未认可的结论" in artifact
        # The domains stay reportable on their own.
        assert "4.0" in artifact
        assert "1.5" in artifact
