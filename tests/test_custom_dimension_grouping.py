from __future__ import annotations

"""Prove the user-defined (non-Scope) category chain: declare dimensions -> compile -> group.

This is the 收尾2 vertical slice at the analysis + service layer. The numbers are fake engine
outputs; what is under test is that a method declaring its own reporting dimension (plant split:
onsite production / purchased energy / upstream materials) compiles through the generic kernel and
re-buckets the engine's per-exchange weighted values into those categories, with no Scope-1-3 code
and no industry branch. A legacy Scope-only method must stay byte-identical.
"""

from nebula_one.analysis.carbon_management import _attach_user_method_grouping
from nebula_one.analysis.method_grouping import group_weighted_by_dimensions
from nebula_one.analysis.pcf import _attach_pcf_user_method_grouping


_PLANT_SPLIT = [
    {
        "id": "plant_split",
        "name": "Plant split",
        "mutually_exclusive": True,
        "requires_full_coverage": False,
        "categories": [
            {"id": "onsite_production", "name": "Onsite production", "is_leaf": True},
            {"id": "purchased_energy", "name": "Purchased energy", "is_leaf": True},
            {"id": "upstream_materials", "name": "Upstream materials", "is_leaf": True},
        ],
    }
]


_SCALED = [
    {"exchange_id": "p1::e_co2", "flow_uuid": "f-co2", "direction": "emission"},
    {"exchange_id": "p1::e_elec", "flow_uuid": "f-elec", "direction": "input"},
    {"exchange_id": "p1::e_steel", "flow_uuid": "f-steel", "direction": "input"},
    {"exchange_id": "p1::e_unclaimed", "flow_uuid": "f-other", "direction": "input"},
]


def _rules() -> list[dict]:
    return [
        {
            "process_instance_id": "p1",
            "exchange_id": "p1::e_co2",
            "direction": "emission",
            "boundary": "include",
            "categories": {"plant_split": "onsite_production"},
        },
        {
            "process_instance_id": "p1",
            "exchange_id": "p1::e_elec",
            "direction": "input",
            "boundary": "include",
            "categories": {"plant_split": "purchased_energy"},
        },
        {
            "process_instance_id": "p1",
            "exchange_id": "p1::e_steel",
            "direction": "input",
            "boundary": "exclude",
            "categories": {},
        },
    ]


_TERMS = [
    {"exchange_id": "p1::e_co2", "contribution_value": 40.0},
    {"exchange_id": "p1::e_elec", "contribution_value": 25.0},
    {"exchange_id": "p1::e_steel", "contribution_value": 10.0},
    {"exchange_id": "p1::e_unclaimed", "contribution_value": 5.0},
]


def _by_id(dim: dict, cat_id: str) -> dict:
    return next(row for row in dim["categories"] if row["id"] == cat_id)


def test_group_weighted_by_dimensions_buckets_custom_categories():
    result = group_weighted_by_dimensions(
        rules=_rules(),
        scaled_exchanges=_SCALED,
        weighted_terms=_TERMS,
        dimensions=_PLANT_SPLIT,
        indicator_unit="kg",
    )
    assert result["reconciled"] is True, result["problems"]
    grouping = result["grouping"]
    dim = grouping["dimensions"]["plant_split"]
    onsite = _by_id(dim, "onsite_production")
    energy = _by_id(dim, "purchased_energy")
    upstream = _by_id(dim, "upstream_materials")

    assert onsite["subtotal"] == 40.0
    assert energy["subtotal"] == 25.0
    # Declared but no contribution assigned: a real zero subtotal would be a fabrication; it stays None.
    assert upstream["subtotal"] is None
    # The excluded steel path moved out of the categories into the excluded total, not deleted.
    assert grouping["excluded"]["value"] == 10.0
    # The unclaimed weighted value is included-but-unclassified: visible, never vanished.
    assert dim["unattributed"]["subtotal"] == 5.0
    # attributed + unattributed equals the known (included) total; the equation never asserts missing=0.
    assert grouping["known_total"] == 40.0 + 25.0 + 5.0


def test_group_weighted_by_dimensions_flags_unknown_category():
    rules = _rules()
    rules[0]["categories"] = {"plant_split": "not_declared"}
    result = group_weighted_by_dimensions(
        rules=rules,
        scaled_exchanges=_SCALED,
        weighted_terms=_TERMS,
        dimensions=_PLANT_SPLIT,
        indicator_unit="kg",
    )
    assert result["reconciled"] is False
    assert any(p["code"] == "GROUPING_UNKNOWN_CATEGORY" for p in result["problems"])


def test_attach_user_method_adds_dimension_grouping_for_custom_method():
    carbon_attribution = {
        "indicator_receipt": {"unit": "kg"},
        "hotspots": [{"exchange_id": t["exchange_id"], "contribution_value": t["contribution_value"]} for t in _TERMS],
    }
    provider_receipt = {
        "user_calculation_method": {
            "method_id": "m1",
            "version": 1,
            "content_sha256": "x",
            "model_revision_id": "r1",
            "execution_plan_sha256": "p",
            "compiler_version": "c",
            "rules": _rules(),
            "reporting_dimensions": _PLANT_SPLIT,
            "execution_plan": {"dimensions": _PLANT_SPLIT},
        }
    }
    out = _attach_user_method_grouping(
        carbon_attribution=carbon_attribution,
        provider_receipt=provider_receipt,
        scaled_elementary_exchanges=_SCALED,
        indicator_key="climate",
    )
    block = out["user_method_grouping"]
    assert block["dimension_grouping"] is not None
    dim = block["dimension_grouping"]["grouping"]["dimensions"]["plant_split"]
    assert _by_id(dim, "onsite_production")["subtotal"] == 40.0
    assert _by_id(dim, "purchased_energy")["subtotal"] == 25.0


def test_attach_user_method_keeps_scope_only_method_shape():
    # A legacy Scope-only method (no declared dimensions) must not gain a dimension_grouping body.
    carbon_attribution = {
        "indicator_receipt": {"unit": "kg"},
        "hotspots": [
            {"exchange_id": "p1::e_co2", "contribution_value": 40.0},
            {"exchange_id": "p1::e_elec", "contribution_value": 25.0},
        ],
    }
    provider_receipt = {
        "user_calculation_method": {
            "method_id": "m2",
            "version": 1,
            "content_sha256": "x",
            "model_revision_id": "r1",
            "execution_plan_sha256": "p",
            "compiler_version": "c",
            "rules": [
                {
                    "process_instance_id": "p1",
                    "exchange_id": "p1::e_co2",
                    "direction": "emission",
                    "boundary_status": "selected",
                    "scope": "scope_1",
                }
            ],
            "reporting_dimensions": [],
        }
    }
    out = _attach_user_method_grouping(
        carbon_attribution=carbon_attribution,
        provider_receipt=provider_receipt,
        scaled_elementary_exchanges=_SCALED,
        indicator_key="climate",
    )
    block = out["user_method_grouping"]
    assert block["dimension_grouping"] is None
    assert block["grouping"]["scopes"]["scope_1"]["value"] == 40.0


def _custom_pcf_receipt() -> dict:
    return {
        "user_calculation_method": {
            "method_id": "m1",
            "version": 1,
            "content_sha256": "x",
            "model_revision_id": "r1",
            "execution_plan_sha256": "p",
            "compiler_version": "c",
            "rules": _rules(),
            "reporting_dimensions": _PLANT_SPLIT,
            "execution_plan": {"dimensions": _PLANT_SPLIT},
        }
    }


def test_attach_pcf_user_method_adds_dimension_grouping():
    """The PCF result page renders a custom (non-Scope) method's categories, not just Scope buckets.

    The frontend ProductCarbonFootprintPanel reads contribution_analysis.user_method_grouping.
    dimension_grouping; a PCF analyze that emits only the Scope path would render nothing for a
    plant-split method, so the PCF grouping path must produce the same dimension rollup as
    carbon-management.
    """
    from nebula_one.analysis.pcf import _attach_pcf_user_method_grouping

    contribution_analysis = {
        "status": "available",
        "hotspots": [
            {"exchange_id": t["exchange_id"], "contribution_value": t["contribution_value"]} for t in _TERMS
        ],
    }
    out = _attach_pcf_user_method_grouping(
        contribution_analysis=contribution_analysis,
        provider_receipt=_custom_pcf_receipt(),
        scaled_elementary_exchanges=_SCALED,
        lcia_results=None,
        indicator_key="climate",
        indicator_unit="kg",
    )
    block = out["user_method_grouping"]
    assert block["grouping_mode"] == "exchange_weighted"
    dim = block["dimension_grouping"]["grouping"]["dimensions"]["plant_split"]
    assert _by_id(dim, "onsite_production")["subtotal"] == 40.0
    assert _by_id(dim, "purchased_energy")["subtotal"] == 25.0
    assert block["dimension_grouping"]["grouping"]["excluded"]["value"] == 10.0
    assert dim["unattributed"]["subtotal"] == 5.0


def test_attach_pcf_user_method_keeps_scope_only_method_shape():
    from nebula_one.analysis.pcf import _attach_pcf_user_method_grouping

    contribution_analysis = {
        "status": "available",
        "hotspots": [{"exchange_id": "p1::e_co2", "contribution_value": 40.0}],
    }
    provider_receipt = {
        "user_calculation_method": {
            "method_id": "m2",
            "version": 1,
            "content_sha256": "x",
            "model_revision_id": "r1",
            "execution_plan_sha256": "p",
            "compiler_version": "c",
            "rules": [
                {
                    "process_instance_id": "p1",
                    "exchange_id": "p1::e_co2",
                    "direction": "emission",
                    "boundary_status": "selected",
                    "scope": "scope_1",
                }
            ],
            "reporting_dimensions": [],
        }
    }
    out = _attach_pcf_user_method_grouping(
        contribution_analysis=contribution_analysis,
        provider_receipt=provider_receipt,
        scaled_elementary_exchanges=_SCALED,
        lcia_results=None,
        indicator_key="climate",
        indicator_unit="kg",
    )
    block = out["user_method_grouping"]
    assert block["dimension_grouping"] is None
    assert block["grouping"]["scopes"]["scope_1"]["value"] == 40.0


# --- export reconciliation (plan R6: page, history, detail and export read one frozen artifact) ---

from io import BytesIO  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from nebula_one.exports.pcf import build_pcf_html, build_pcf_xlsx  # noqa: E402


def _dimension_grouping_block() -> dict:
    return {
        "grouping_mode": "exchange_weighted",
        "execution_status": "completed",
        "counts_as_complete_pcf": False,
        "method": {
            "method_id": "m1", "version": 1, "content_sha256": "x",
            "execution_plan_sha256": "p", "compiler_version": "c",
        },
        "dimension_grouping": {
            "reconciled": True,
            "grouping": {
                "indicator_unit": "kg",
                "known_total": 70.0,
                "known_term_count": 3,
                "excluded": {"value": 10.0, "unit": "kg", "term_count": 1, "missing_count": 0},
                "dimensions": {
                    "plant_split": {
                        "categories": [
                            {"id": "onsite_production", "name": "Onsite production",
                             "subtotal": 40.0, "unit": "kg", "term_count": 1, "missing_count": 0},
                            {"id": "purchased_energy", "name": "Purchased energy",
                             "subtotal": 25.0, "unit": "kg", "term_count": 1, "missing_count": 0},
                            {"id": "upstream_materials", "name": "Upstream materials",
                             "subtotal": None, "unit": "kg", "term_count": 0, "missing_count": 1},
                        ],
                        "unattributed": {"subtotal": 5.0, "term_count": 1},
                    }
                },
            },
        },
    }


def _pcf_result(umg: dict | None) -> dict:
    analysis = {"status": "available", "coverage": "complete", "hotspots": [],
                "user_method_grouping": umg}
    return {
        "schema_version": "pcf.product.v1",
        "content_sha256": "result-hash",
        "total_footprint": {"indicator_key": "climate", "value": 70.0, "unit": "kg"},
        "functional_unit": {"amount": 1, "unit": "kg", "reference_flow_uuid": "f-product"},
        "boundary_claim": {"claim_scope": "cradle-to-gate", "coverage": "complete",
                           "complete_product_declaration": False},
        "source_snapshot": {"id": "s1", "content_sha256": "solve-hash",
                            "system_revision_id": "r1", "model_graph_sha256": "g1"},
        "method_package": {"lcia_method": "EF v3.1", "reference": "pcf.product.ef31",
                           "content_sha256": "pkg-hash"},
        "component_results": [],
        "contribution_analysis": analysis,
        "audit": {"status": "ok", "checks": [], "limitations": []},
        "provenance": {},
        "data_quality": {"database_release": "tiangong-1"},
    }


def test_xlsx_user_method_sheet_renders_dimension_subtotals():
    stream = build_pcf_xlsx(_pcf_result(_dimension_grouping_block()))
    workbook = load_workbook(BytesIO(stream), read_only=True)
    assert "用户方法归组" in workbook.sheetnames
    cells = [c.value for row in workbook["用户方法归组"].iter_rows() for c in row]
    assert "40.0" in cells and "25.0" in cells
    assert "对账一致" in cells


def test_xlsx_no_method_export_keeps_historical_sheet_set():
    stream = build_pcf_xlsx(_pcf_result(None))
    assert load_workbook(BytesIO(stream), read_only=True).sheetnames == [
        "核算摘要", "LCIA结果", "贡献热点", "审计与追溯",
    ]


def test_html_report_renders_user_method_grouping_section():
    html_text = build_pcf_html(_pcf_result(_dimension_grouping_block()), "<svg></svg>")
    assert "用户方法归组" in html_text
    assert "40.0" in html_text and "25.0" in html_text
    assert "对账一致" in html_text


def test_html_report_without_method_omits_grouping_section():
    html_text = build_pcf_html(_pcf_result(None), "<svg></svg>")
    assert "用户方法归组" not in html_text

