from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from nebula_one.hashing import canonical_json_sha256
from nebula_one.models import AnalysisRun


METHOD_PACKAGE = Path("examples/pcf-product-gold-standard/method-package.json")


def _indicator(key: str, value: float, unit: str = "kg CO2-Eq") -> dict:
    return {
        "method_en": "EF v3.1",
        "indicator_en": "Climate change - GWP100",
        "canonical_indicator_key": key,
        "value": value,
        "unit": unit,
        "indicator_unit": unit,
    }


def _prepare_snapshot(
    client,
    *,
    elementary_receipts=None,
    total_unit="kg CO2-Eq",
    include_contribution_receipts=False,
    mapped_receipts=None,
    provenance_extra=None,
    with_factor_gap=False,
):
    package_payload = json.loads(METHOD_PACKAGE.read_text(encoding="utf-8"))
    package = client.post("/api/v1/method-packages", json=package_payload)
    assert package.status_code == 201, package.text
    package = package.json()

    application = client.post(
        "/api/v1/applications",
        json={"name": "PCF gold standard", "application_type": "system_analysis"},
    ).json()
    model = client.post(
        f"/api/v1/applications/{application['id']}/system-models",
        json={
            "name": "PCF gold model",
            "provider": "nebula_lca",
            "upstream_project_id": "project-1",
            "upstream_version": 3,
        },
    ).json()
    revision_id = model["upstream_receipt"]["base_revision_id"]
    revision = client.get(f"/api/v1/system-revisions/{revision_id}").json()
    run = AnalysisRun(
        system_model_id=model["id"],
        system_revision_id=revision_id,
        capability_key="solve.nebula.v1",
        idempotency_key="pcf-gold-solve",
        request_sha256="a" * 64,
        parameters={
            "demand": [
                {
                    "reference_flow_uuid": "flow-product",
                    "amount": 1.0,
                    "unit": "kg",
                    "unit_group_uuid": "mass",
                    "unit_group_version": "1",
                }
            ],
            "lcia_methods": ["EF v3.1"],
        },
        status="succeeded",
        model_graph_sha256=model["graph_sha256"],
        provider_receipt={},
        result={},
    )
    with client.app.state.container.database.session() as db:
        db.add(run)
        db.commit()
        db.refresh(run)
        run_id = run.id
    elementary_flow_receipts = (
        elementary_receipts
        if elementary_receipts is not None
        else [
            {
                "exchange_id": "process::co2-out",
                "source_namespace": "tiangong_open_data",
                "flow_uuid": "08a91e70-3ddc-11dd-923d-0050c2490048",
                "version": "03.00.004",
                "method": "EF v3.1",
                "factor_hash": "b" * 64,
            }
        ]
    )
    scaled_elementary_exchanges = [
        {
                "process_id": "process",
                "exchange_id": "process::co2-out",
                "flow_uuid": "08a91e70-3ddc-11dd-923d-0050c2490048",
                "amount": 2.0,
                "unit": "kg",
                "unit_group": "mass",
                "direction": "output",
                "exchange_type": "elementary",
            }
    ] if elementary_flow_receipts else []
    lcia_results = {
        "schema_version": "provider.lcia.response.v1",
        "method": "EF v3.1",
        "database_release": "EF3.1",
        "inventory_hash": "d" * 64,
        "indicator_results": [
            _indicator("climate change", 2.0, total_unit),
            _indicator("climate change: fossil", 2.0),
            _indicator("climate change: biogenic", 0.0),
            _indicator("climate change: land use and land use change", 0.0),
        ],
        "runtime_assets": [{"runtime_id": "ef31-gold"}],
    }
    if include_contribution_receipts:
        solve_hash = "e" * 64
        runtime_assets_hash = "f" * 64
        term = {
            "schema_version": "provider.lcia.contribution-term.v1",
            "exchange_id": "process::co2-out",
            "process_uuid": "process",
            "source_namespace": "tiangong_open_data",
            "flow_uuid": "08a91e70-3ddc-11dd-923d-0050c2490048",
            "flow_version": "03.00.004",
            "amount": 2.0,
            "amount_unit": "kg",
            "direction": "output",
            "contribution_value": 2.0,
            "contribution_unit": "kg CO2-Eq",
            "cf_method": "EF v3.1",
            "cf_database_release": "EF3.1",
            "cf_indicator_key": "climate change",
            "cf_runtime_assets_hash": runtime_assets_hash,
            "solve_hash": solve_hash,
            "cf_hash": "1" * 64,
        }
        receipt = {
            "schema_version": "provider.lcia.indicator-contribution-receipt.v1",
            "method": "EF v3.1",
            "database_release": "EF3.1",
            "canonical_indicator_key": "climate change",
            "indicator_unit": "kg CO2-Eq",
            "indicator_value": 2.0,
            "contribution_total": 2.0,
            "reconciliation_delta": 0.0,
            "terms": [term],
            "solve_hash": solve_hash,
            "provider_commit": "provider-test",
            "runtime_assets_hash": runtime_assets_hash,
        }
        receipt["content_hash"] = canonical_json_sha256(receipt)
        lcia_results.update(
            {
                "solve_hash": solve_hash,
                "provider_commit": "provider-test",
                "runtime_assets_hash": runtime_assets_hash,
                "indicator_contribution_receipts": [receipt],
            }
        )
        lcia_results["contribution_receipts_hash"] = canonical_json_sha256([receipt])
    if with_factor_gap:
        assert include_contribution_receipts
        gap = {"exchange_id": "process::missing", "process_uuid": "process",
               "flow_uuid": "missing-cf-flow", "flow_version": "1", "amount": 0.1,
               "amount_unit": "kg", "status": "factor_not_found", "numeric_contribution": None}
        scaled_elementary_exchanges.append({"exchange_id": gap["exchange_id"], "process_id": "process",
            "flow_uuid": gap["flow_uuid"], "amount": 0.1, "unit": "kg", "direction": "output",
            "exchange_type": "elementary", "unit_group": "mass"})
        elementary_flow_receipts.append({"exchange_id": gap["exchange_id"], "flow_uuid": gap["flow_uuid"],
            "version": "1", "factor_count": 0, "factor_hash": canonical_json_sha256([])})
        receipt["characterization_coverage"] = {
            "schema_version": "provider.lcia.characterization-coverage.v1", "status": "partial",
            "value_kind": "known_subtotal", "uncharacterized_exchanges": [gap],
            "runtime_assets_hash": runtime_assets_hash,
        }
        receipt["content_hash"] = canonical_json_sha256({k: v for k, v in receipt.items() if k != "content_hash"})
        lcia_results["contribution_receipts_hash"] = canonical_json_sha256([receipt])
    snapshot_payload = {
        "system_revision_id": revision_id,
        "provider": "nebula_lca",
        "provider_receipt": {
            "endpoint": "/api/provider/v1/solve",
            "status": "completed",
            "upstream_run_id": "pcf-gold-provider-run",
            "system_revision_hash": revision["content_sha256"],
            "analysis_run_id": run_id,
            "database_release": "EF3.1",
            "solver_version": "gold-standard",
            "elementary_flow_receipts": elementary_flow_receipts,
            "background_process_receipts": [],
            "mapped_background_process_receipts": mapped_receipts or [],
            "provenance": {"request_hash": "c" * 64, **(provenance_extra or {})},
        },
        "functional_unit": revision["graph_snapshot"]["functionalUnit"],
        "database_release": "EF3.1",
        "solver_version": "gold-standard",
        "activity_vector": [],
        "scaled_technosphere_exchanges": [
            {
                "process_id": "process",
                "exchange_id": "process::product-out",
                "flow_uuid": "flow-product",
                "amount": 1.0,
                "unit": "kg",
                "unit_group": "mass",
                "direction": "output",
                "exchange_type": "technosphere",
            }
        ],
        "scaled_elementary_exchanges": scaled_elementary_exchanges,
        "internal_transfers": [],
        "boundary_flows": [],
        "process_balance_residuals": [],
        "contribution_graph": {},
        "lcia_results": lcia_results,
    }
    snapshot = client.post(
        f"/api/v1/applications/{application['id']}/solved-system-snapshots",
        json=snapshot_payload,
    )
    assert snapshot.status_code == 201, snapshot.text
    return snapshot.json(), f"{package['package_id']}@{package['version']}"


def test_mapped_background_is_partial_coverage_not_a_complete_declaration(client):
    snapshot, package_ref = _prepare_snapshot(client, mapped_receipts=[{
        "process_uuid": "synthetic-provider", "claim_role": "partial_background_leaf",
    }])
    response = client.post(
        f"/api/v1/solved-system-snapshots/{snapshot['id']}/product-carbon-footprint",
        json={"method_package_ref": package_ref},
    )
    assert response.status_code == 200, response.text
    claim = response.json()["boundary_claim"]
    assert claim["claim_scope"] == "model_defined_partial_background"
    assert claim["coverage"] == "foreground_with_pinned_background"
    assert claim["complete_product_declaration"] is False
    report = client.get(
        f"/api/v1/solved-system-snapshots/{snapshot['id']}/product-carbon-footprint/report.html",
        params={"method_package_ref": package_ref},
    )
    assert report.status_code == 200
    assert "model_defined_partial_background" in report.text


def test_partial_characterization_is_visible_in_http_and_exports(client) -> None:
    snapshot, package_ref = _prepare_snapshot(client, include_contribution_receipts=True, with_factor_gap=True)
    path = f"/api/v1/solved-system-snapshots/{snapshot['id']}/product-carbon-footprint"
    response = client.post(path, json={"method_package_ref": package_ref})
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["total_footprint"]["value"] == 2.0
    assert result["total_footprint"]["value_kind"] == "known_subtotal"
    assert result["data_quality"]["characterization_coverage"]["status"] == "partial"
    assert result["data_quality"]["characterization_coverage"]["uncharacterized_exchanges"][0]["numeric_contribution"] is None
    assert result["boundary_claim"]["complete_product_declaration"] is False
    exported = client.get(path + "/export.xlsx", params={"method_package_ref": package_ref})
    assert exported.status_code == 200
    book = load_workbook(BytesIO(exported.content))
    assert any("已匹配贡献小计" in str(row[0]) for row in book.worksheets[0].values)
    report = client.get(path + "/report.html", params={"method_package_ref": package_ref})
    assert report.status_code == 200
    assert "已匹配贡献小计" in report.text


def test_product_pcf_gold_standard_is_deterministic(client) -> None:
    snapshot, package_ref = _prepare_snapshot(client)
    path = (
        f"/api/v1/solved-system-snapshots/{snapshot['id']}"
        "/product-carbon-footprint"
    )
    first = client.post(path, json={"method_package_ref": package_ref})
    replay = client.post(path, json={"method_package_ref": package_ref})

    assert first.status_code == replay.status_code == 200
    assert first.json() == replay.json()
    result = first.json()
    assert result["total_footprint"] == {
        "value": 2.0,
        "unit": "kg CO2-Eq",
        "indicator_key": "climate change",
    }
    assert result["functional_unit"]["amount"] == 1.0
    assert result["boundary_claim"] == {
        "claim_scope": "foreground_boundary",
        "coverage": "foreground_only",
        "complete_product_declaration": False,
    }
    assert result["audit"]["status"] == "passed"
    assert result["contribution_analysis"]["status"] == "unavailable"
    content = dict(result)
    content.pop("content_sha256")
    assert result["content_sha256"] == canonical_json_sha256(content)
    chart = client.post(
        f"{path}/chart.svg", json={"method_package_ref": package_ref}
    )
    assert chart.status_code == 200
    assert chart.headers["content-type"].startswith("image/svg+xml")
    assert "Product carbon footprint" in chart.text
    assert "2 kg CO2-Eq" in chart.text
    chart_get = client.get(
        f"{path}/chart.svg", params={"method_package_ref": package_ref}
    )
    assert chart_get.status_code == 200
    assert chart_get.text == chart.text


def test_unit_product_projection_reuses_the_solve_and_rescales_the_whole_result(client) -> None:
    snapshot, package_ref = _prepare_snapshot(client)
    path = f"/api/v1/solved-system-snapshots/{snapshot['id']}/product-carbon-footprint"

    # total basis (default) is the historical identity: no result_basis envelope is invented.
    total = client.post(path, json={"method_package_ref": package_ref}).json()
    assert total["total_footprint"]["value"] == 2.0 and "result_basis" not in total
    solve_hash = total["source_snapshot"]["content_sha256"]

    # per_reference_output over a 1 t base for a 1 kg reference output: T * 1000 / 1.
    intensity = client.post(
        path,
        json={"method_package_ref": package_ref, "result_basis": "per_reference_output",
              "base_amount": 1, "base_unit": "t"},
    )
    assert intensity.status_code == 200, intensity.text
    body = intensity.json()
    assert body["total_footprint"]["value"] == 2000.0
    assert body["total_footprint"]["unit"] == "kg CO2-Eq/1 t"
    # the same frozen solve is reused, never re-run: solve hash is byte-identical
    assert body["source_snapshot"]["content_sha256"] == solve_hash
    assert body["result_basis"]["source"]["solve_hash"] == solve_hash
    assert body["result_basis"]["source"]["raw_total"] == {"value": 2.0, "unit": "kg CO2-Eq"}
    # the seam that broke silently: the foreground HTTP path must resolve the bound reference flow
    # against the snapshot's actual producing outputs and freeze the unique hit's identities.
    reference = body["result_basis"]["reference_output"]
    assert reference["selector"] == "reference_flow_uuid" and reference["value"] == "flow-product"
    assert reference["producing_process_uuid"] == "process"
    assert reference["producing_exchange_id"] == "process::product-out"
    assert reference["identity"] == "process::product-out"
    # the projected product re-hashes over its own body so the API is self-consistent
    unsigned = dict(body)
    unsigned.pop("content_sha256")
    assert body["content_sha256"] == canonical_json_sha256(unsigned)

    # the headline follows the basis: intensity is per-product, the default is the raw total.
    report_intensity = client.get(
        f"{path}/report.html",
        params={"method_package_ref": package_ref, "result_basis": "per_reference_output",
                "base_amount": 1, "base_unit": "t"},
    )
    assert report_intensity.status_code == 200
    assert "单位产品气候变化" in report_intensity.text
    assert "2000" in report_intensity.text
    report_total = client.get(
        f"{path}/report.html", params={"method_package_ref": package_ref}
    )
    assert "单位产品气候变化" not in report_total.text
    assert "本次核算总量气候变化" in report_total.text

    # an intensity export carries its own ETag distinct from the total export.
    xlsx_total = client.get(f"{path}/export.xlsx", params={"method_package_ref": package_ref})
    xlsx_intensity = client.get(
        f"{path}/export.xlsx",
        params={"method_package_ref": package_ref, "result_basis": "per_reference_output",
                "base_amount": 1, "base_unit": "t"},
    )
    assert xlsx_total.status_code == xlsx_intensity.status_code == 200
    assert xlsx_total.headers["ETag"] != xlsx_intensity.headers["ETag"]


def test_intensity_projection_requires_an_exact_base(client) -> None:
    snapshot, package_ref = _prepare_snapshot(client)
    path = f"/api/v1/solved-system-snapshots/{snapshot['id']}/product-carbon-footprint"
    response = client.post(
        path, json={"method_package_ref": package_ref, "result_basis": "per_reference_output"},
    )
    assert response.status_code == 422
    bad_unit = client.post(
        path, json={"method_package_ref": package_ref, "result_basis": "per_reference_output",
                    "base_amount": 1, "base_unit": "kWh"},
    )
    assert bad_unit.status_code == 422


def test_product_pcf_fails_closed_without_elementary_receipts(client) -> None:
    snapshot, package_ref = _prepare_snapshot(client, elementary_receipts=[])
    response = client.post(
        f"/api/v1/solved-system-snapshots/{snapshot['id']}/product-carbon-footprint",
        json={"method_package_ref": package_ref},
    )
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "PCF_ELEMENTARY_FLOW_RECEIPTS_REQUIRED"


def test_product_pcf_fails_closed_on_impact_unit_mismatch(client) -> None:
    snapshot, package_ref = _prepare_snapshot(client, total_unit="kg")
    response = client.post(
        f"/api/v1/solved-system-snapshots/{snapshot['id']}/product-carbon-footprint",
        json={"method_package_ref": package_ref},
    )
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "PCF_RESULT_UNIT_MISMATCH"


def test_product_pcf_consumes_exact_contribution_receipt(client) -> None:
    snapshot, package_ref = _prepare_snapshot(
        client, include_contribution_receipts=True
    )
    response = client.post(
        f"/api/v1/solved-system-snapshots/{snapshot['id']}/product-carbon-footprint",
        json={"method_package_ref": package_ref},
    )
    assert response.status_code == 200, response.text
    analysis = response.json()["contribution_analysis"]
    assert analysis["status"] == "available"
    assert analysis["coverage"] == "exact_exchange_contributions"
    assert analysis["hotspots"][0]["exchange_id"] == "process::co2-out"
    assert analysis["hotspots"][0]["share_of_total"] == 1.0


def test_product_pcf_exports_excel_and_printable_html(client) -> None:
    snapshot, package_ref = _prepare_snapshot(
        client, include_contribution_receipts=True
    )
    base = f"/api/v1/solved-system-snapshots/{snapshot['id']}/product-carbon-footprint"
    result = client.post(base, json={"method_package_ref": package_ref}).json()
    params = {"method_package_ref": package_ref}

    workbook_response = client.get(f"{base}/export.xlsx", params=params)
    assert workbook_response.status_code == 200
    assert workbook_response.headers["x-nebula-result-sha256"] == result["content_sha256"]
    workbook = load_workbook(BytesIO(workbook_response.content), read_only=True)
    assert workbook.sheetnames == ["核算摘要", "LCIA结果", "贡献热点", "审计与追溯"]
    assert workbook["核算摘要"]["B3"].value == result["content_sha256"]
    assert workbook["贡献热点"]["C2"].value == "process::co2-out"

    html_response = client.get(f"{base}/report.html", params=params)
    assert html_response.status_code == 200
    assert html_response.headers["x-nebula-result-sha256"] == result["content_sha256"]
    assert "产品碳足迹技术核算报告" in html_response.text
    assert "window.print()" in html_response.text
    assert result["content_sha256"] in html_response.text


MAPPING_DEPENDENCIES = [{
    "package_id": "synthetic-elementary-mapping-core", "package_version": "1.0.0",
    "package_hash": "9" * 64, "role": "certified_identity_adjudication",
}]


def test_elementary_mapping_dependencies_survive_freeze_read_and_export(client):
    snapshot, package_ref = _prepare_snapshot(
        client, provenance_extra={"elementary_mapping_dependencies": MAPPING_DEPENDENCIES}
    )
    base = f"/api/v1/solved-system-snapshots/{snapshot['id']}/product-carbon-footprint"
    params = {"method_package_ref": package_ref}
    result = client.post(base, json=params).json()
    assert result["provenance"]["elementary_mapping_dependencies"] == MAPPING_DEPENDENCIES

    workbook_response = client.get(f"{base}/export.xlsx", params=params)
    assert workbook_response.headers["x-nebula-result-sha256"] == result["content_sha256"]
    workbook = load_workbook(BytesIO(workbook_response.content), read_only=True)
    audit = {row[0].value: row[1].value for row in workbook["审计与追溯"].iter_rows(min_row=2, max_col=2)}
    assert json.loads(audit["Provider provenance"])["elementary_mapping_dependencies"] == MAPPING_DEPENDENCIES

    report = client.get(f"{base}/report.html", params=params)
    assert "基本流映射依赖包" in report.text
    assert "synthetic-elementary-mapping-core" in report.text


def test_snapshot_without_the_dependency_contract_gains_no_invented_record(client):
    snapshot, package_ref = _prepare_snapshot(client)
    base = f"/api/v1/solved-system-snapshots/{snapshot['id']}/product-carbon-footprint"
    params = {"method_package_ref": package_ref}
    result = client.post(base, json=params).json()
    assert "elementary_mapping_dependencies" not in result["provenance"]
    assert "基本流映射依赖包" not in client.get(f"{base}/report.html", params=params).text
    content = dict(result)
    content.pop("content_sha256")
    assert result["content_sha256"] == canonical_json_sha256(content)
