from copy import deepcopy
from io import BytesIO

import pytest
from openpyxl import load_workbook

from nebula_one.analysis.frozen_report import cbam_report, organization_report
from nebula_one.errors import DomainError
from nebula_one.exports.frozen_report import export_frozen_report_xlsx, render_frozen_report_html


def graph():
    return {'nodes': [{'id': 'plant', 'process_uuid': 'process', 'data': {'name': 'Steel', 'outputs': [],
                     'emissions': [{'id': 'co2', 'name': 'Carbon dioxide'}, {'id': 'nox', 'name': 'Nitrogen oxides'}]}}]}


def cbam():
    return {'attributed_emissions': {'direct': {'value': '0.003', 'unit': 't CO2e', 'source_value': '3', 'source_unit': 'kg CO2e',
             'exchanges': [{'exchange_id': 'plant::co2', 'amount': '3', 'unit': 'kg CO2e'}]}, 'indirect': None},
            'specific_embedded_emissions': {'value': None}, 'coverage': {'exchanges': [
                {'exchange_id': 'plant::co2', 'process_uuid': 'process', 'amount': '8', 'unit': 'kg', 'disposition': 'counted'},
                {'exchange_id': 'plant::nox', 'process_uuid': 'process', 'amount': '2', 'unit': 'kg', 'disposition': 'excluded_with_basis'}],
                'open_items': [{'item': 'Pending electricity factor', 'state': 'missing_data'}]}}


def test_cbam_uses_recorded_emissions_and_not_raw_quantity_or_tonne_summary():
    report = cbam_report(cbam(), graph(), {'fields': []})
    assert report['details'][0]['quantity'] == '8'
    assert report['details'][0]['contribution_value'] == '3'
    assert report['process_contributions'][0]['contribution_value'] == 3
    assert report['summary'][0]['value'] == '0.003'
    assert report['details'][1]['contribution_value'] is None
    assert report['complete_total'] is None


def test_contribution_mismatch_refuses_instead_of_inventing_a_reconciliation():
    result = cbam()
    result['attributed_emissions']['direct']['source_value'] = '4'
    with pytest.raises(DomainError, match='Frozen detail contributions'):
        cbam_report(result, graph(), {'fields': []})


def test_organization_scales_once_and_distinguishes_absent_and_explicit_zero_cf():
    terms = [{'process_uuid': 'process', 'exchange_id': 'plant::co2', 'amount': 8, 'amount_unit': 'kg', 'contribution_value': 3, 'contribution_unit': 'kg CO2-Eq', 'cf_presence': 'explicit_nonzero'},
             {'process_uuid': 'process', 'exchange_id': 'plant::nox', 'amount': 2, 'amount_unit': 'kg', 'contribution_value': 0, 'contribution_unit': 'kg CO2-Eq', 'cf_presence': 'implicit_zero'},
             {'process_uuid': 'process', 'exchange_id': 'plant::explicit-zero', 'amount': 4, 'amount_unit': 'kg', 'contribution_value': 0, 'contribution_unit': 'kg CO2-Eq', 'cf_presence': 'explicit_zero'}]
    result = {'activity_basis': {'scale_factor': 2}, 'organizational_boundary': {'controlled_process_ids': ['plant']},
              'groups': {'on_site_direct': {'value': 6, 'unit': 'kg CO2-Eq', 'status': 'calculated'},
                         'purchased_energy_lifecycle': {'value': 0, 'unit': 'kg CO2-Eq', 'status': 'calculated'},
                         'other_value_chain': {'value': None, 'unit': 'kg CO2-Eq', 'status': 'partial', 'categories': [
                             {'category': 'missing', 'name': 'Missing supplier', 'status': 'missing_activity'},
                             {'category': 'outside', 'name': 'Downstream', 'status': 'not_in_boundary'}]}},
              'combined_total': None, 'coverage': {'missing_categories': ['missing']},
              'climate_source': {'indicator_key': 'climate', 'uncharacterized_terms': [{'exchange_id': 'plant::nox'}]}}
    before = deepcopy(result)
    report = organization_report(result, {'lcia_results': {'indicator_contribution_receipts': [{'canonical_indicator_key': 'climate', 'terms': terms}]}}, graph(), {'fields': []})
    assert report['process_contributions'][0]['contribution_value'] == 6
    assert report['details'][0]['quantity'] == 8
    assert report['uncharacterized'][0]['contribution_value'] is None
    assert report['details'][2]['status'] == 'explicit_zero_cf'
    assert report['details'][2]['contribution_value'] == 0
    assert [r['name'] for r in report['gaps']] == ['Missing supplier']
    assert [r['name'] for r in report['outside_boundary']] == ['Downstream']
    assert result == before


def test_full_exports_keep_every_row_nulls_names_and_exact_source_in_appendix():
    report = cbam_report(cbam(), graph(), {'fields': []})
    report.update(title='钢铁报告', snapshot_id='frozen-snapshot', revision_id='frozen-revision', method_version='1.0.0',
                  basis={'source_value': '1000', 'source_unit': 'kg'}, demand=[], inputs=[
                      {'process': '钢铁', 'label': f'流{i}', 'group': 'input', 'value': i, 'unit': 'kg', 'key': f'exact-exchange-{i}'} for i in range(28)],
                  sources={'content_sha256': 'a' * 64, 'all_rows': list(range(9000))})
    book = load_workbook(BytesIO(export_frozen_report_xlsx(report)))
    assert book['冻结输入'].max_row == 29
    assert book['核算明细']['F3'].value is None
    assert book['核算明细']['F2'].value == '3'
    assert book['方法及来源'].max_row > 9000
    html = render_frozen_report_html(report)
    body, appendix = html.split('<details><summary>方法及来源附录</summary>')
    assert '流27' in body and '钢铁报告' in body
    assert 'exact-exchange-' not in body
    assert 'frozen-revision' in appendix and 'a' * 64 in appendix
    assert '完整合计未计算' in body
