"""The reportable face of a frozen native background climate answer.

Nebula LCA answers the pinned climate contract with per-exchange terms whose emission factors are
licensed content, so an artifact may carry only what the contract says is publishable: values,
units, exact version references and coverage statements. Every row written here is selected field by
field from the validated answer, and no term list is ever serialised.
"""
from __future__ import annotations

import html
from io import BytesIO
from typing import Any

from openpyxl import Workbook

from ..mapped_background_schemas import (
    ProviderNativeBackgroundClimateRequest,
    ProviderNativeBackgroundClimateResponse,
)
from .pcf import _append_key_values, _style_sheet, _table, _text

MERGE_STATUS_LABELS = {
    "merged": "来源已合并为合计",
    "split_diagnostic": "来源未合并，按域分别给出",
}
COVERAGE_LABELS = {
    "explicit_nonzero": "有非零排放因子",
    "explicit_zero": "因子明确为零",
    "no_cf": "未纳入特征化",
}
DOMAIN_LABELS = {
    "foreground": "前台生产过程",
    "native_background": "已关联的背景供应",
}
STATEMENT_LABELS = {
    "not_assessed": "未评估",
    "not_asserted_by_local_vector_metadata": "本地向量元数据未作声明",
    "partial_background_leaf_only_not_complete_product_carbon_footprint": (
        "背景覆盖止于已关联的背景供应叶子，本结果不构成完整产品碳足迹"
    ),
}


def _label(value: Any, labels: dict[str, str]) -> str:
    text = _text(value)
    return labels.get(text, text)


def _coverage_text(coverage: dict[str, int]) -> str:
    if not coverage:
        return "来源未返回覆盖计数"
    return " · ".join(
        f"{_label(state, COVERAGE_LABELS)} {count} 项"
        for state, count in sorted(coverage.items())
    )


def _split_blockers(answer: ProviderNativeBackgroundClimateResponse) -> list[str]:
    # The source's own key order is kept so an artifact reads exactly like the on-screen answer.
    return [
        " · ".join(f"{key}: {_text(value)}" for key, value in blocker.items())
        if isinstance(blocker, dict)
        else _text(blocker)
        for blocker in answer.merge_blockers
    ]


def _summary_rows(
    answer: ProviderNativeBackgroundClimateResponse,
    request: ProviderNativeBackgroundClimateRequest,
    basis: str,
) -> list[tuple[str, Any]]:
    indicator = answer.indicator
    total = answer.unified_climate_total
    merged = answer.merge_status == "merged"
    return [
        ("结果口径", _label(answer.claim_limit, STATEMENT_LABELS)),
        ("指标", f"{indicator.name} · {indicator.category}"),
        ("核算方法", indicator.method),
        ("合计状态", _label(answer.merge_status, MERGE_STATUS_LABELS)),
        (
            "气候变化合计（含已关联背景）"
            if merged
            else "气候变化合计（来源未合并，因此无值）",
            f"{_text(total)} {indicator.unit}" if merged else "—",
        ),
        ("前台生产过程", f'{_text(answer.domains["foreground"].contribution_value)} {indicator.unit}'),
        (
            "已关联的背景供应",
            f'{_text(answer.domains["native_background"].contribution_value)} {indicator.unit}',
        ),
        ("结果基准", basis),
        ("背景数据版本", request.database_release),
    ]


def _domain_rows(answer: ProviderNativeBackgroundClimateResponse) -> list[list[Any]]:
    return [
        [
            _label(key, DOMAIN_LABELS),
            domain.source_namespace,
            domain.contribution_value,
            domain.indicator_unit,
            len(domain.terms),
            _coverage_text(domain.coverage),
            domain.negative_cf_exchange_count,
        ]
        for key, domain in answer.domains.items()
    ]


def _statement_rows(answer: ProviderNativeBackgroundClimateResponse) -> list[tuple[str, Any]]:
    negative_terms = sum(
        domain.negative_cf_exchange_count for domain in answer.domains.values()
    )
    rows = [
        (
            "负贡献说明",
            f"含 {negative_terms} 项负贡献交换，代表抵消性的投入输出，不等同于永久碳移除"
            if negative_terms
            else "本次答案没有负贡献交换",
        ),
        (
            "碳储存永久性",
            _label(answer.carbon_storage_permanence, STATEMENT_LABELS),
        ),
        (
            "背景系统模型",
            _label(answer.provider_system_model_evidence, STATEMENT_LABELS),
        ),
    ]
    if answer.merge_status == "merged":
        rows.append(("分域数值处理", "两域数值由来源合并为一个合计，无需自行相加"))
    else:
        rows.append(
            ("分域数值处理", "各域口径不一致，把两域数值相加会得到一个来源未认可的结论")
        )
        for index, blocker in enumerate(_split_blockers(answer), start=1):
            rows.append((f"未合并原因 {index}", blocker))
    rows.append(
        (
            "逐条排放因子",
            "来源按许可数据返回逐交换明细，本文件只保留数值、单位、版本引用与覆盖说明",
        )
    )
    return rows


def _version_rows(
    answer: ProviderNativeBackgroundClimateResponse,
    request: ProviderNativeBackgroundClimateRequest,
    source_contract: dict[str, Any],
) -> list[tuple[str, Any]]:
    return [
        ("方法包标识", answer.method_package_id),
        ("方法包版本", answer.method_package_version),
        ("方法包内容哈希", answer.method_package_sha256),
        ("运行时清单哈希", answer.runtime_manifest_sha256),
        ("请求钉住的方法包哈希", request.method_package_sha256),
        ("请求钉住的运行时清单哈希", request.runtime_manifest_sha256),
        ("流快照哈希", request.flow_snapshot_hash),
        ("参考依赖快照哈希", request.reference_dependency_snapshot_hash),
        ("来源合同仓库", source_contract.get("repository")),
        ("来源合同提交", source_contract.get("commit")),
        ("来源合同文档", source_contract.get("document")),
        ("来源合同文档哈希", source_contract.get("document_sha256")),
    ]


def build_native_climate_xlsx(
    envelope: dict[str, Any],
    answer: ProviderNativeBackgroundClimateResponse,
    request: ProviderNativeBackgroundClimateRequest,
    *,
    snapshot_id: str,
    snapshot_sha256: str,
    basis: str,
) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "气候核算摘要"
    summary.append(["项目", "值"])
    _append_key_values(summary, _summary_rows(answer, request, basis))
    _append_key_values(
        summary,
        [
            ("结果 Schema", envelope.get("schema_version")),
            ("求解快照 ID", snapshot_id),
            ("求解快照内容哈希", snapshot_sha256),
        ],
    )

    domains = workbook.create_sheet("来源分域与覆盖")
    domains.append(
        [
            "领域",
            "来源命名空间",
            "贡献值",
            "指标单位",
            "交换项数",
            "覆盖说明",
            "负贡献交换数",
        ]
    )
    for row in _domain_rows(answer):
        domains.append(row)

    statements = workbook.create_sheet("口径与限制")
    statements.append(["项目", "说明"])
    _append_key_values(statements, _statement_rows(answer))

    versions = workbook.create_sheet("版本与来源证据")
    versions.append(["项目", "值"])
    _append_key_values(versions, _version_rows(answer, request, envelope["source_contract"]))

    for sheet in workbook.worksheets:
        _style_sheet(sheet)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _facts(rows: list[tuple[str, Any]]) -> str:
    return "".join(
        f'<div class="fact"><span>{html.escape(key)}</span>{html.escape(_text(value))}</div>'
        for key, value in rows
    )


def build_native_climate_html(
    envelope: dict[str, Any],
    answer: ProviderNativeBackgroundClimateResponse,
    request: ProviderNativeBackgroundClimateRequest,
    *,
    snapshot_id: str,
    snapshot_sha256: str,
    basis: str,
) -> str:
    merged = answer.merge_status == "merged"
    indicator = answer.indicator
    total = answer.unified_climate_total
    headline = (
        f"{html.escape(_text(total))} {html.escape(indicator.unit)}"
        if merged
        else "来源未合并，因此无合计值"
    )
    notice = (
        '<p class="notice">本结果覆盖至已关联的背景供应叶子，不构成完整产品碳足迹；两域数值已由来源合并，无需自行相加。</p>'
        if merged
        else '<p class="notice">各域口径不一致，来源未给出合计；把下列分域数值相加会得到一个来源未认可的结论。</p>'
    )
    return f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Nebula One 部分背景下的气候变化结果</title>
<style>
:root{{--ink:#17324a;--muted:#60758a;--line:#d9e3ea;--soft:#f4f8fb;--accent:#1976a3}}
*{{box-sizing:border-box}} body{{margin:0;background:#edf3f7;color:var(--ink);font:14px/1.55 "Segoe UI","Microsoft YaHei",sans-serif}}
.page{{max-width:1100px;margin:24px auto;background:white;padding:40px 48px;box-shadow:0 8px 32px #15324b18}}
.toolbar{{display:flex;justify-content:flex-end;margin-bottom:12px}} button{{height:40px;padding:0 18px;border:0;border-radius:8px;background:var(--accent);color:white;font-weight:600;cursor:pointer}}
h1{{font-size:26px;margin:0 0 4px}} h2{{font-size:18px;margin:30px 0 12px;border-bottom:1px solid var(--line);padding-bottom:8px}}
.subtitle{{color:var(--muted)}} .metric{{margin:24px 0;padding:20px 24px;background:var(--soft);border-left:4px solid var(--accent)}}
.metric strong{{display:block;font-size:30px}} .facts{{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:10px 24px}}
.fact{{border-bottom:1px solid var(--line);padding:8px 0;word-break:break-all}} .fact span{{display:block;color:var(--muted);font-size:12px}}
table{{width:100%;border-collapse:collapse;font-size:12px}} th,td{{border:1px solid var(--line);padding:8px;text-align:left;vertical-align:top}} th{{background:var(--soft)}}
.notice{{padding:12px 16px;background:#fff7e8;border:1px solid #f1d49a;border-radius:8px}} code{{word-break:break-all;font-size:11px}}
@media(max-width:720px){{.page{{margin:0;padding:24px 18px;box-shadow:none}}.facts{{grid-template-columns:1fr}}.table-wrap{{overflow:auto}}}}
@media print{{body{{background:white}}.page{{max-width:none;margin:0;padding:12mm;box-shadow:none}}.toolbar{{display:none}}h2{{break-after:avoid}}table,.metric,.notice{{break-inside:avoid}}}}
</style></head><body><main class="page">
<div class="toolbar"><button type="button" onclick="window.print()">打印 / 另存 PDF</button></div>
<h1>部分背景下的气候变化结果</h1>
<div class="subtitle">Nebula One · {html.escape(answer.method_package_id)} {html.escape(answer.method_package_version)}</div>
<section class="metric"><span>气候变化合计（含已关联背景）</span><strong>{headline}</strong><small>{html.escape(indicator.name)} · {html.escape(indicator.category)} · {html.escape(indicator.method)} · 结果基准 {html.escape(basis)}</small></section>
{notice}
<section><h2>核算口径</h2><div class="facts">{_facts(_summary_rows(answer, request, basis))}</div></section>
<section><h2>来源分域与覆盖</h2><div class="table-wrap">{_table(['领域', '来源命名空间', '贡献值', '指标单位', '交换项数', '覆盖说明', '负贡献交换数'], _domain_rows(answer))}</div></section>
<section><h2>口径与限制</h2><div class="facts">{_facts(_statement_rows(answer))}</div></section>
<section><h2>版本与来源证据</h2><div class="facts">{_facts([*(_version_rows(answer, request, envelope['source_contract'])), ('求解快照 ID', snapshot_id), ('求解快照内容哈希', snapshot_sha256)])}</div></section>
</main></body></html>"""
