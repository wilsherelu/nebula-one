from __future__ import annotations

import html
import json
from io import BytesIO
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


_HEADER_FILL = PatternFill("solid", fgColor="15324B")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_SECTION_FILL = PatternFill("solid", fgColor="DCEAF2")


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    if isinstance(value, bool):
        return "是" if value else "否"
    return str(value)


def _basis_heading(result: dict[str, Any]) -> str:
    # The headline must follow the actual projected basis, never a fixed word: `total` is the
    # whole-demand solve total, only `per_reference_output` is a per-product intensity.
    partial = (result.get("data_quality", {}).get("characterization_coverage") or {}).get("status") == "partial"
    prefix = "已匹配贡献小计 · " if partial else ""
    basis = (result.get("result_basis") or {}).get("basis")
    if basis == "per_reference_output":
        base = result["result_basis"]
        return prefix + f"单位产品气候变化（每 {base.get('base_amount')} {base.get('base_unit')}）"
    return prefix + "本次核算总量气候变化"


def _append_key_values(sheet, rows: Iterable[tuple[str, Any]]) -> None:
    for key, value in rows:
        sheet.append([key, _text(value)])


def _style_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for column in range(1, sheet.max_column + 1):
        width = max(
            12,
            min(
                52,
                max(len(_text(sheet.cell(row=row, column=column).value)) for row in range(1, sheet.max_row + 1)) + 2,
            ),
        )
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def build_pcf_xlsx(result: dict[str, Any]) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "核算摘要"
    summary.append(["项目", "值"])
    total = result["total_footprint"]
    functional_unit = result["functional_unit"]
    boundary = result["boundary_claim"]
    source = result["source_snapshot"]
    method = result["method_package"]
    _append_key_values(
        summary,
        [
            ("结果 Schema", result["schema_version"]),
            ("结果内容哈希", result["content_sha256"]),
            ("产品 Flow UUID", functional_unit.get("reference_flow_uuid")),
            ("功能单位", f'{functional_unit.get("amount")} {functional_unit.get("unit")}'),
            (_basis_heading(result), f'{total.get("value")} {total.get("unit")}'),
            ("LCIA 方法", method.get("lcia_method")),
            ("方法包", method.get("reference")),
            ("方法包哈希", method.get("content_sha256")),
            ("核算边界", boundary.get("claim_scope")),
            ("数据覆盖", boundary.get("coverage")),
            ("完整产品声明", boundary.get("complete_product_declaration")),
            ("求解快照", source.get("id")),
            ("求解快照哈希", source.get("content_sha256")),
            ("System revision", source.get("system_revision_id")),
            ("模型图哈希", source.get("model_graph_sha256")),
        ],
    )

    basis = result.get("result_basis")
    characterization = result.get("data_quality", {}).get("characterization_coverage")
    if characterization is not None:
        _append_key_values(summary, [("特征化覆盖状态", characterization["status"]),
                                    ("未匹配因子（未计入小计，非零排放声明）", characterization["uncharacterized_exchanges"])])
    if basis:
        raw_total = (basis.get("source") or {}).get("raw_total") or {}
        _append_key_values(summary, [
            ("结果口径", basis.get("basis")),
            ("展示基准", f'{basis.get("base_amount")} {basis.get("base_unit")}'),
            ("换算基准哈希", basis.get("basis_sha256")),
            ("换算比率（基准/参考产出）", basis.get("conversion_ratio_base_over_reference")),
            ("原始求解总量", f'{raw_total.get("value")} {raw_total.get("unit")}'),
        ])

    components = workbook.create_sheet("LCIA结果")
    components.append(["指标", "数值", "单位"])
    components.append([total.get("indicator_key"), total.get("value"), total.get("unit")])
    for item in result.get("component_results") or []:
        components.append([item.get("indicator_key"), item.get("value"), item.get("unit")])

    hotspots = workbook.create_sheet("贡献热点")
    hotspots.append(
        [
            "排名",
            "Process ID",
            "Exchange ID",
            "Flow UUID",
            "Flow 版本",
            "清单数量",
            "清单单位",
            "贡献值",
            "贡献单位",
            "占总量比例",
            "CF 哈希",
        ]
    )
    analysis = result.get("contribution_analysis") or {}
    for item in analysis.get("hotspots") or []:
        identity = item.get("flow_identity") or {}
        hotspots.append(
            [
                item.get("rank"),
                item.get("process_id"),
                item.get("exchange_id"),
                identity.get("flow_uuid"),
                identity.get("version"),
                item.get("amount"),
                item.get("amount_unit"),
                item.get("contribution_value"),
                item.get("contribution_unit"),
                item.get("share_of_total"),
                item.get("cf_hash"),
            ]
        )
    if hotspots.max_row == 1:
        hotspots.append(["不可用", analysis.get("reason")])

    audit = workbook.create_sheet("审计与追溯")
    audit.append(["项目", "值"])
    _append_key_values(
        audit,
        [
            ("审计状态", result.get("audit", {}).get("status")),
            ("审计检查", result.get("audit", {}).get("checks")),
            ("限制条件", result.get("audit", {}).get("limitations")),
            ("贡献分析状态", analysis.get("status")),
            ("贡献分析覆盖", analysis.get("coverage")),
            ("贡献凭据", analysis.get("indicator_receipt")),
            ("数据质量", result.get("data_quality")),
            ("Provider provenance", result.get("provenance")),
        ],
    )

    # A cited user method's grouping is part of the frozen result; the export must read the same
    # values the page shows (plan R6: 结果、历史、详情和导出读取同一冻结产物). Absent a method the
    # sheet is omitted so a no-method PCF export keeps its historical sheet set byte-for-byte.
    umg = analysis.get("user_method_grouping")
    if umg:
        _append_user_method_sheet(workbook, umg)

    for sheet in workbook.worksheets:
        _style_sheet(sheet)
    summary["A1"].fill = _SECTION_FILL
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _append_user_method_sheet(workbook: "Workbook", umg: dict[str, Any]) -> None:
    sheet = workbook.create_sheet("用户方法归组")
    method = umg.get("method") or {}
    sheet.append(["项目", "值"])
    for label, value in (
        ("方法 ID", method.get("method_id")),
        ("方法版本", method.get("version")),
        ("内容哈希", method.get("content_sha256")),
        ("执行计划哈希", method.get("execution_plan_sha256")),
        ("编译器版本", method.get("compiler_version")),
        ("归组模式", umg.get("grouping_mode")),
        ("执行状态", umg.get("execution_status")),
        ("计入完整 PCF", umg.get("counts_as_complete_pcf", False)),
    ):
        sheet.append([label, _text(value)])
    dg = umg.get("dimension_grouping")
    if dg:
        aggregate = dg.get("grouping") or {}
        unit = aggregate.get("indicator_unit")
        sheet.append([])
        sheet.append(["维度", "类别 ID", "类别名称", "小计", "单位", "项数", "缺失项"])
        for dim_id, dim in (aggregate.get("dimensions") or {}).items():
            for cat in dim.get("categories") or []:
                sheet.append([
                    dim_id, cat.get("id"), cat.get("name"),
                    _text(cat.get("subtotal")), unit,
                    cat.get("term_count"), cat.get("missing_count"),
                ])
            unattributed = dim.get("unattributed") or {}
            sheet.append([
                dim_id, "__unattributed__", "未归集",
                _text(unattributed.get("subtotal")), unit,
                unattributed.get("term_count"), "",
            ])
        excluded = aggregate.get("excluded") or {}
        sheet.append([])
        sheet.append([
            "明确排除", "", "", _text(excluded.get("value")), unit,
            excluded.get("term_count"), excluded.get("missing_count"),
        ])
        sheet.append(["已知贡献合计", "", "", _text(aggregate.get("known_total")), unit, aggregate.get("known_term_count"), ""])
        sheet.append(["对账一致", "", "", _text(dg.get("reconciled")), "", "", ""])
        return
    grouping = umg.get("grouping") or {}
    scopes = grouping.get("scopes") or {}
    unit = grouping.get("indicator_unit")
    sheet.append([])
    sheet.append(["范围", "数值", "单位", "项数"])
    for scope_key in ("scope_1", "scope_2", "scope_3"):
        entry = scopes.get(scope_key) or {}
        sheet.append([scope_key, _text(entry.get("value")), entry.get("unit"), entry.get("term_count")])
    excluded = grouping.get("excluded") or {}
    sheet.append(["明确排除", _text(excluded.get("value")), excluded.get("unit"), excluded.get("term_count")])
    sheet.append(["贡献合计", _text(grouping.get("contribution_total")), unit, ""])
    sheet.append(["对账一致", _text(grouping.get("reconciled")), "", ""])


def _table(headers: list[str], rows: list[list[Any]]) -> str:
    head = "".join(f"<th>{html.escape(header)}</th>" for header in headers)
    body = "".join(
        "<tr>" + "".join(f"<td>{html.escape(_text(value))}</td>" for value in row) + "</tr>"
        for row in rows
    )
    return f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"


def _user_method_grouping_html(umg: dict[str, Any] | None) -> str:
    if not umg:
        return ""
    method = umg.get("method") or {}
    facts = "".join(
        f'<div class="fact"><span>{html.escape(label)}</span>{html.escape(_text(value))}</div>'
        for label, value in (
            ("方法 ID", method.get("method_id")),
            ("方法版本", method.get("version")),
            ("内容哈希", method.get("content_sha256")),
            ("执行计划哈希", method.get("execution_plan_sha256")),
            ("归组模式", umg.get("grouping_mode")),
            ("执行状态", umg.get("execution_status")),
        )
    )
    dg = umg.get("dimension_grouping")
    if dg:
        aggregate = dg.get("grouping") or {}
        unit = aggregate.get("indicator_unit")
        rows: list[list[Any]] = []
        for dim_id, dim in (aggregate.get("dimensions") or {}).items():
            for cat in dim.get("categories") or []:
                rows.append([
                    dim_id, cat.get("id"), cat.get("name"),
                    _text(cat.get("subtotal")), unit, cat.get("term_count"), cat.get("missing_count"),
                ])
            unattributed = dim.get("unattributed") or {}
            rows.append([dim_id, "__unattributed__", "未归集", _text(unattributed.get("subtotal")), unit, unattributed.get("term_count"), ""])
        excluded = aggregate.get("excluded") or {}
        rows.append(["", "excluded", "明确排除", _text(excluded.get("value")), unit, excluded.get("term_count"), excluded.get("missing_count")])
        rows.append(["", "known_total", "已知贡献合计", _text(aggregate.get("known_total")), unit, aggregate.get("known_term_count"), ""])
        table = _table(["维度", "类别 ID", "类别名称", "小计", "单位", "项数", "缺失项"], rows)
        reconciled = _text(dg.get("reconciled"))
    else:
        grouping = umg.get("grouping") or {}
        scopes = grouping.get("scopes") or {}
        rows = [
            [key, _text((scopes.get(key) or {}).get("value")), (scopes.get(key) or {}).get("unit"), (scopes.get(key) or {}).get("term_count")]
            for key in ("scope_1", "scope_2", "scope_3")
        ]
        excluded = grouping.get("excluded") or {}
        rows.append(["明确排除", _text(excluded.get("value")), excluded.get("unit"), excluded.get("term_count")])
        rows.append(["贡献合计", _text(grouping.get("contribution_total")), grouping.get("indicator_unit"), ""])
        table = _table(["范围", "数值", "单位", "项数"], rows)
        reconciled = _text(grouping.get("reconciled"))
    return (
        '<section><h2>用户方法归组</h2>'
        f'<p>所用方法版本：{html.escape(_text(method.get("version")))}；'
        f'计入完整 PCF：{html.escape(_text(umg.get("counts_as_complete_pcf", False)))}；'
        f'对账一致：{html.escape(reconciled)}</p>'
        f'<div class="facts">{facts}</div>'
        f'<div class="table-wrap">{table}</div></section>'
    )


def build_pcf_html(result: dict[str, Any], chart_svg: str) -> str:
    total = result["total_footprint"]
    functional_unit = result["functional_unit"]
    boundary = result["boundary_claim"]
    analysis = result.get("contribution_analysis") or {}
    component_rows = [
        [item.get("indicator_key"), item.get("value"), item.get("unit")]
        for item in result.get("component_results") or []
    ]
    hotspot_rows = []
    for item in analysis.get("hotspots") or []:
        identity = item.get("flow_identity") or {}
        share = item.get("share_of_total")
        hotspot_rows.append(
            [
                item.get("rank"),
                item.get("process_id"),
                item.get("exchange_id"),
                f'{identity.get("flow_uuid", "")}@{identity.get("version", "")}',
                item.get("contribution_value"),
                item.get("contribution_unit"),
                "" if share is None else f"{float(share) * 100:.2f}%",
            ]
        )
    if not hotspot_rows:
        hotspot_rows = [["—", analysis.get("reason") or "无贡献热点数据", "", "", "", "", ""]]
    limitations = result.get("audit", {}).get("limitations") or []
    limitations_html = "".join(f"<li>{html.escape(_text(item))}</li>" for item in limitations)
    dependencies = result.get("provenance", {}).get("elementary_mapping_dependencies")
    dependencies_html = (
        f'<div class="fact"><span>基本流映射依赖包</span><code>{html.escape(_text(dependencies))}</code></div>'
        if dependencies else ""
    )
    grouping_html = _user_method_grouping_html(analysis.get("user_method_grouping"))
    return f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Nebula One 产品碳足迹技术核算报告</title>
<style>
:root{{--ink:#17324a;--muted:#60758a;--line:#d9e3ea;--soft:#f4f8fb;--accent:#1976a3}}
*{{box-sizing:border-box}} body{{margin:0;background:#edf3f7;color:var(--ink);font:14px/1.55 "Segoe UI","Microsoft YaHei",sans-serif}}
.page{{max-width:1100px;margin:24px auto;background:white;padding:40px 48px;box-shadow:0 8px 32px #15324b18}}
.toolbar{{display:flex;justify-content:flex-end;margin-bottom:12px}} button{{height:40px;padding:0 18px;border:0;border-radius:8px;background:var(--accent);color:white;font-weight:600;cursor:pointer}}
h1{{font-size:28px;margin:0 0 4px}} h2{{font-size:18px;margin:30px 0 12px;border-bottom:1px solid var(--line);padding-bottom:8px}}
.subtitle{{color:var(--muted)}} .metric{{margin:24px 0;padding:20px 24px;background:var(--soft);border-left:4px solid var(--accent)}}
.metric strong{{display:block;font-size:30px}} .facts{{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:10px 24px}}
.fact{{border-bottom:1px solid var(--line);padding:8px 0}} .fact span{{display:block;color:var(--muted);font-size:12px}}
table{{width:100%;border-collapse:collapse;font-size:12px}} th,td{{border:1px solid var(--line);padding:8px;text-align:left;vertical-align:top}} th{{background:var(--soft)}}
.chart svg{{width:100%;height:auto}} code{{word-break:break-all;font-size:11px}} .notice{{padding:12px 16px;background:#fff7e8;border:1px solid #f1d49a;border-radius:8px}}
@media(max-width:720px){{.page{{margin:0;padding:24px 18px;box-shadow:none}}.facts{{grid-template-columns:1fr}}.table-wrap{{overflow:auto}}}}
@media print{{body{{background:white}}.page{{max-width:none;margin:0;padding:12mm;box-shadow:none}}.toolbar{{display:none}}h2{{break-after:avoid}}table,.metric,.notice{{break-inside:avoid}}}}
</style></head><body><main class="page">
<div class="toolbar"><button type="button" onclick="window.print()">打印 / 另存 PDF</button></div>
<h1>产品碳足迹技术核算报告</h1><div class="subtitle">Nebula One · {html.escape(_text(result['method_package'].get('reference')))}</div>
<section class="metric"><span>{html.escape(_basis_heading(result))}</span><strong>{html.escape(_text(total.get('value')))} {html.escape(_text(total.get('unit')))}</strong><small>功能单位：{html.escape(_text(functional_unit.get('amount')))} {html.escape(_text(functional_unit.get('unit')))}</small></section>
<section><h2>核算对象与边界</h2><div class="facts">
<div class="fact"><span>产品 Flow</span>{html.escape(_text(functional_unit.get('reference_flow_uuid')))}</div>
<div class="fact"><span>LCIA 方法</span>{html.escape(_text(result['method_package'].get('lcia_method')))}</div>
<div class="fact"><span>核算边界</span>{html.escape(_text(boundary.get('claim_scope')))}</div>
<div class="fact"><span>数据覆盖</span>{html.escape(_text(boundary.get('coverage')))}</div></div>
<p class="notice">本报告按冻结模型的实际求解边界披露结果，不构成第三方核证、认证产品声明或法规申报。</p></section>
<section class="chart"><h2>气候变化结果</h2>{chart_svg}</section>
<section><h2>LCIA 分项</h2><div class="table-wrap">{_table(['指标','数值','单位'], component_rows)}</div></section>
<section><h2>贡献热点</h2><p>状态：{html.escape(_text(analysis.get('status')))}；覆盖：{html.escape(_text(analysis.get('coverage')))}</p><div class="table-wrap">{_table(['排名','Process','Exchange','Flow','贡献值','单位','占总量'], hotspot_rows)}</div></section>
{('<section><h2>未匹配评价因子</h2><p>下列流保留在清单中，未计入已匹配贡献小计；缺因子不代表零排放。</p>' + _table(['流 UUID', '数量', '单位'], [[g.get('flow_uuid'), g.get('amount'), g.get('amount_unit')] for g in result['data_quality']['characterization_coverage']['uncharacterized_exchanges']]) + '</section>') if (result.get('data_quality', {}).get('characterization_coverage') or {}).get('status') == 'partial' else ''}{grouping_html}
<section><h2>限制条件</h2><ul>{limitations_html}</ul></section>
<section><h2>审计与追溯</h2><div class="facts">
<div class="fact"><span>结果哈希</span><code>{html.escape(_text(result.get('content_sha256')))}</code></div>
<div class="fact"><span>求解快照哈希</span><code>{html.escape(_text(result['source_snapshot'].get('content_sha256')))}</code></div>
<div class="fact"><span>模型图哈希</span><code>{html.escape(_text(result['source_snapshot'].get('model_graph_sha256')))}</code></div>
<div class="fact"><span>数据库版本</span>{html.escape(_text(result.get('data_quality', {}).get('database_release')))}</div>{dependencies_html}</div></section>
</main></body></html>"""
