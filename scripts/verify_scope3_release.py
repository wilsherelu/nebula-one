"""Verify the committed standalone payload with fresh storage and real browser actions."""
from __future__ import annotations
import argparse
import hashlib
import json
from pathlib import Path
import subprocess
import sys

from openpyxl import load_workbook
import e3_packaged_runtime_verify as P
from live_scope3_rules_acceptance import prepare
from live_official_rule_equivalence import run as verify_library


def verify(installation, data, evidence):
    evidence.mkdir(parents=True,exist_ok=True)
    data.mkdir(parents=True,exist_ok=True)
    lca=one=editor=None
    report={'installation':str(installation),'data':str(data),'checks':[]}
    def check(name,ok,detail=None):
        report['checks'].append(dict(name=name,passed=bool(ok),detail=detail))
        if not ok: raise AssertionError(name)
    try:
        runtime=installation/'resources/runtime'
        manifest=json.loads((runtime/'payload-manifest.json').read_text(encoding='utf-8'))
        report['one_commit']=manifest['one_commit'];report['lca_commit']=manifest['lca_commit']
        mismatches=[]
        for row in manifest['files']:
            file=runtime/row['path']
            if not file.is_file() or hashlib.sha256(file.read_bytes()).hexdigest()!=row['sha256']: mismatches.append(row['path'])
        check('every_payload_file_matches_manifest',not mismatches,{'files':len(manifest['files']),'mismatches':mismatches})
        lca_port,one_port=P._free_port(),P._free_port()
        lca_url,one_url=f'http://127.0.0.1:{lca_port}',f'http://127.0.0.1:{one_port}'
        lca,_=P._launch_lca(installation,data,lca_port)
        check('lca_ready',P._wait_ready(lca_url+'/health','release LCA'))
        editor=P._EditorServer(runtime/'editor').start()
        one=P._launch_one(runtime,data,one_port,lca_url,editor.origin)
        check('one_ready',P._wait_ready(one_url+'/api/v1/health','release One'))
        check('one_version_0_1_1',P._get(one_url+'/api/v1/health')['version']=='0.1.1')
        library = verify_library(one_url, lca_url, evidence / 'template-library')
        check('packaged_template_library_and_rule_equivalence', library['failed'] == 0,
              {'passed': library['passed'], 'report': library['output']})
        fixture=prepare(one_url,lca_url,evidence)
        ui=evidence/'ui'
        with (evidence/'ui.log').open('w',encoding='utf-8') as log:
            status=subprocess.run(['node',str(P.REPO/'scripts/ui_scope3_rules_acceptance.mjs'),
                str(evidence/'scope3-fixture.json'),str(ui)],cwd=P.REPO,stdout=log,stderr=subprocess.STDOUT).returncode
        check('real_browser_scope3_journey',status==0,{'log':str(evidence/'ui.log')})
        results=json.loads((ui/'scope3-ui-summary.json').read_text(encoding='utf-8'))
        check('every_ui_assertion_passes',all(row['pass'] for row in results['checks']),{'checks':len(results['checks'])})
        phase=results['phases']['regrouped']
        book=load_workbook(next((ui/'downloads').glob('*.xlsx')),data_only=True)
        summary={row[0]:row[1] for row in book['摘要'].iter_rows(min_row=2,values_only=True)}
        check('downloaded_excel_scope3',abs(summary['其他上下游排放']-fixture['independent_scope3'])<1e-7)
        groups={row[0]:row[1] for row in book['供应链分组'].iter_rows(min_row=2,values_only=True)}
        check('downloaded_excel_custom_group',abs(groups['材料组 B']-fixture['independent_scope3'])<1e-7)
        html=next((ui/'downloads').glob('*.html')).read_text(encoding='utf-8')
        check('downloaded_html_same_method',fixture['method_id'] in html and phase['org']['content_sha256'] in html)
        records={p['record']:p['org']['content_sha256'] for p in results['phases'].values()}
        pid=one.pid;one.terminate();one.wait(timeout=20)
        one=P._launch_one(runtime,data,one_port,lca_url,editor.origin)
        check('one_restarted_same_storage',P._wait_ready(one_url+'/api/v1/health','release One restarted'))
        for record,sha in records.items():
            read=P._get(f'{one_url}/api/v1/applications/{fixture["application_id"]}/organization-lifecycle-inventories/{record}')
            check('persisted_record_'+record,read['content_sha256']==sha)
        report['restart_readback']=records;report['initial_one_pid']=pid
        report['status']='passed'
    except Exception as error:
        report['status']='failed';report['error']=str(error)
        raise
    finally:
        for proc in (one,lca):
            if proc and proc.poll() is None: proc.terminate();proc.wait(timeout=20)
        if editor: editor.stop()
        (evidence/'release-scope3.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'status':report['status'],'checks':len(report['checks'])}))


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--installation',type=Path,required=True)
    parser.add_argument('--data',type=Path,required=True)
    parser.add_argument('--evidence',type=Path,required=True)
    args=parser.parse_args();verify(args.installation.resolve(),args.data.resolve(),args.evidence.resolve())
